from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    UploadFile,
    File,
)
from openpyxl import load_workbook
from pydantic import BaseModel, EmailStr
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.connection import get_db
from app.models.hod import HOD
from app.models.line_manager import LineManager
from app.models.target_audience import TargetAudience


router = APIRouter(
    prefix="/api/recipients",
    tags=["Recipients"],
)


# =========================================================
# HELPERS
# =========================================================

async def get_target_audience_by_name(
    db: AsyncSession,
    name: str,
):
    """
    Find a TargetAudience by name, case-insensitively.
    """
    if not name:
        return None

    result = await db.execute(
        select(TargetAudience).where(
            func.lower(TargetAudience.name) == name.strip().lower()
        )
    )

    return result.scalar_one_or_none()


async def get_or_create_target_audience(
    db: AsyncSession,
    name: str,
):
    """
    Resolve a target audience from the master table.

    A new audience is created only when the Settings Excel
    explicitly contains an audience that does not yet exist.
    """
    cleaned_name = (name or "").strip()

    if not cleaned_name:
        return None

    audience = await get_target_audience_by_name(
        db,
        cleaned_name,
    )

    if audience:
        return audience

    audience = TargetAudience(
        name=cleaned_name,
    )

    db.add(audience)
    await db.flush()

    return audience


def normalize_recipient_type(value):
    value = str(value or "").strip().upper()

    aliases = {
        "HOD": "HOD",
        "HEAD OF DEPARTMENT": "HOD",
        "LINE_MANAGER": "LINE_MANAGER",
        "LINE MANAGER": "LINE_MANAGER",
        "MANAGER": "LINE_MANAGER",
    }

    return aliases.get(value, value)


# =========================================================
# REQUEST SCHEMAS
# =========================================================

class RecipientUpdate(BaseModel):
    name: str
    email: EmailStr
    whatsapp_number: str | None = None
    target_audience: str


# =========================================================
# HOD API
# =========================================================

@router.get("/hods")
async def get_hods(
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(HOD).order_by(HOD.id)
    )

    hods = result.scalars().all()

    records = []

    for hod in hods:
        audience = None

        if getattr(hod, "audience_id", None):
            audience_result = await db.execute(
                select(TargetAudience).where(
                    TargetAudience.id == hod.audience_id
                )
            )
            audience = audience_result.scalar_one_or_none()

        records.append(
            {
                "id": hod.id,
                "name": hod.name,
                "email": hod.email,
                "whatsapp_number": hod.whatsapp_number,
                "department_id": getattr(
                    hod,
                    "department_id",
                    None,
                ),
                "audience_id": getattr(
                    hod,
                    "audience_id",
                    None,
                ),
                "target_audience": (
                    audience.name
                    if audience
                    else None
                ),
                "active": getattr(
                    hod,
                    "active",
                    True,
                ),
            }
        )

    return records


# =========================================================
# LINE MANAGER API
# =========================================================

@router.get("/line-managers")
async def get_line_managers(
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(LineManager).order_by(LineManager.id)
    )

    managers = result.scalars().all()

    records = []

    for manager in managers:
        audience = None

        if getattr(manager, "audience_id", None):
            audience_result = await db.execute(
                select(TargetAudience).where(
                    TargetAudience.id == manager.audience_id
                )
            )
            audience = audience_result.scalar_one_or_none()

        records.append(
            {
                "id": manager.id,
                "name": manager.name,
                "email": manager.email,
                "whatsapp_number": manager.whatsapp_number,
                "department_id": getattr(
                    manager,
                    "department_id",
                    None,
                ),
                "audience_id": getattr(
                    manager,
                    "audience_id",
                    None,
                ),
                "target_audience": (
                    audience.name
                    if audience
                    else None
                ),
                "active": getattr(
                    manager,
                    "active",
                    True,
                ),
            }
        )

    return records


# =========================================================
# TARGET AUDIENCE MASTER API
# =========================================================

@router.get("/audiences")
async def get_target_audiences(
    db: AsyncSession = Depends(get_db),
):
    """
    Return all target audiences.

    The Settings UI can use this endpoint to populate a
    Target Audience dropdown.
    """
    result = await db.execute(
        select(TargetAudience).order_by(
            TargetAudience.name
        )
    )

    audiences = result.scalars().all()

    return [
        {
            "id": audience.id,
            "name": audience.name,
        }
        for audience in audiences
    ]


# =========================================================
# MASTER DATA
# =========================================================

@router.get("/master")
async def get_master_recipients(
    db: AsyncSession = Depends(get_db),
):
    """
    Return all HODs and Line Managers in one combined list.

    Recipient assignment is based on TargetAudience.id.
    """

    hod_result = await db.execute(
        select(HOD).order_by(HOD.id)
    )

    manager_result = await db.execute(
        select(LineManager).order_by(LineManager.id)
    )

    audience_result = await db.execute(
        select(TargetAudience)
    )

    hods = hod_result.scalars().all()
    managers = manager_result.scalars().all()
    audiences = audience_result.scalars().all()

    audience_map = {
        audience.id: audience.name
        for audience in audiences
    }

    records = []

    # -----------------------------------------------------
    # HODs
    # -----------------------------------------------------

    for hod in hods:
        records.append(
            {
                "id": hod.id,
                "type": "HOD",
                "name": hod.name,
                "email": hod.email,
                "whatsapp_number": hod.whatsapp_number,
                "department_id": getattr(
                    hod,
                    "department_id",
                    None,
                ),
                "audience_id": getattr(
                    hod,
                    "audience_id",
                    None,
                ),
                "target_audience": audience_map.get(
                    getattr(hod, "audience_id", None)
                ),
                "active": getattr(
                    hod,
                    "active",
                    True,
                ),
            }
        )

    # -----------------------------------------------------
    # Line Managers
    # -----------------------------------------------------

    for manager in managers:
        records.append(
            {
                "id": manager.id,
                "type": "LINE_MANAGER",
                "name": manager.name,
                "email": manager.email,
                "whatsapp_number": manager.whatsapp_number,
                "department_id": getattr(
                    manager,
                    "department_id",
                    None,
                ),
                "audience_id": getattr(
                    manager,
                    "audience_id",
                    None,
                ),
                "target_audience": audience_map.get(
                    getattr(manager, "audience_id", None)
                ),
                "active": getattr(
                    manager,
                    "active",
                    True,
                ),
            }
        )

    return records


# =========================================================
# UPDATE HOD
# =========================================================

@router.put("/hods/{hod_id}")
async def update_hod(
    hod_id: int,
    data: RecipientUpdate,
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(HOD).where(
            HOD.id == hod_id
        )
    )

    hod = result.scalar_one_or_none()

    if not hod:
        raise HTTPException(
            status_code=404,
            detail="HOD not found",
        )

    audience = await get_target_audience_by_name(
        db,
        data.target_audience,
    )

    if not audience:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Target Audience '{data.target_audience}' "
                "does not exist. Add it through Settings "
                "master data first."
            ),
        )

    hod.name = data.name.strip()
    hod.email = str(data.email).strip()
    hod.whatsapp_number = (
        data.whatsapp_number.strip()
        if data.whatsapp_number
        else None
    )
    hod.audience_id = audience.id

    await db.commit()
    await db.refresh(hod)

    return {
        "success": True,
        "message": "HOD updated successfully",
        "data": {
            "id": hod.id,
            "type": "HOD",
            "name": hod.name,
            "email": hod.email,
            "whatsapp_number": hod.whatsapp_number,
            "audience_id": hod.audience_id,
            "target_audience": audience.name,
            "department_id": getattr(
                hod,
                "department_id",
                None,
            ),
            "active": getattr(
                hod,
                "active",
                True,
            ),
        },
    }


# =========================================================
# UPDATE LINE MANAGER
# =========================================================

@router.put("/line-managers/{manager_id}")
async def update_line_manager(
    manager_id: int,
    data: RecipientUpdate,
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(LineManager).where(
            LineManager.id == manager_id
        )
    )

    manager = result.scalar_one_or_none()

    if not manager:
        raise HTTPException(
            status_code=404,
            detail="Line Manager not found",
        )

    audience = await get_target_audience_by_name(
        db,
        data.target_audience,
    )

    if not audience:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Target Audience '{data.target_audience}' "
                "does not exist. Add it through Settings "
                "master data first."
            ),
        )

    manager.name = data.name.strip()
    manager.email = str(data.email).strip()
    manager.whatsapp_number = (
        data.whatsapp_number.strip()
        if data.whatsapp_number
        else None
    )
    manager.audience_id = audience.id

    await db.commit()
    await db.refresh(manager)

    return {
        "success": True,
        "message": "Line Manager updated successfully",
        "data": {
            "id": manager.id,
            "type": "LINE_MANAGER",
            "name": manager.name,
            "email": manager.email,
            "whatsapp_number": manager.whatsapp_number,
            "audience_id": manager.audience_id,
            "target_audience": audience.name,
            "department_id": getattr(
                manager,
                "department_id",
                None,
            ),
            "active": getattr(
                manager,
                "active",
                True,
            ),
        },
    }


# =========================================================
# DELETE HOD
# =========================================================

@router.delete("/hods/{hod_id}")
async def delete_hod(
    hod_id: int,
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(HOD).where(
            HOD.id == hod_id
        )
    )

    hod = result.scalar_one_or_none()

    if not hod:
        raise HTTPException(
            status_code=404,
            detail="HOD not found",
        )

    await db.delete(hod)
    await db.commit()

    return {
        "success": True,
        "message": "HOD deleted successfully",
    }


# =========================================================
# DELETE LINE MANAGER
# =========================================================

@router.delete("/line-managers/{manager_id}")
async def delete_line_manager(
    manager_id: int,
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(LineManager).where(
            LineManager.id == manager_id
        )
    )

    manager = result.scalar_one_or_none()

    if not manager:
        raise HTTPException(
            status_code=404,
            detail="Line Manager not found",
        )

    await db.delete(manager)
    await db.commit()

    return {
        "success": True,
        "message": "Line Manager deleted successfully",
    }


# =========================================================
# MASTER DATA EXCEL UPLOAD
# =========================================================

@router.post("/master/upload")
async def upload_master_recipients(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    """
    Upload HOD / Line Manager master data from Excel.

    Expected columns:

        Type
        Name
        Email
        WhatsApp Number
        Target Audience

    Target Audience is resolved to:

        target_audiences.id

    and stored in:

        HOD.audience_id
        LineManager.audience_id
    """

    # -----------------------------------------------------
    # FILE VALIDATION
    # -----------------------------------------------------

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected",
        )

    filename = file.filename.lower()

    if not filename.endswith(
        (".xlsx", ".xlsm")
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Only .xlsx or .xlsm Excel files "
                "are supported"
            ),
        )

    try:
        contents = await file.read()

        workbook = load_workbook(
            filename=BytesIO(contents),
            data_only=True,
        )

        worksheet = workbook.active

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Unable to read Excel file: {str(exc)}"
            ),
        )

    # -----------------------------------------------------
    # READ ROWS
    # -----------------------------------------------------

    rows = list(
        worksheet.iter_rows(
            values_only=True
        )
    )

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="Excel file is empty",
        )

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    headers = [
        str(value).strip()
        if value is not None
        else ""
        for value in rows[0]
    ]

    required_columns = [
        "Type",
        "Name",
        "Email",
        "WhatsApp Number",
        "Target Audience",
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in headers
    ]

    if missing_columns:
        raise HTTPException(
            status_code=400,
            detail={
                "message": (
                    "Required columns are missing"
                ),
                "missing_columns": missing_columns,
                "required_columns": required_columns,
            },
        )

    column_indexes = {
        column: headers.index(column)
        for column in headers
    }

    # -----------------------------------------------------
    # RESULT TRACKING
    # -----------------------------------------------------

    created = []
    updated = []
    errors = []

    # -----------------------------------------------------
    # HELPER
    # -----------------------------------------------------

    def get_value(row, column_name):
        index = column_indexes.get(
            column_name
        )

        if index is None:
            return None

        if index >= len(row):
            return None

        value = row[index]

        if value is None:
            return None

        return str(value).strip()

    # -----------------------------------------------------
    # AUDIENCE CACHE
    # -----------------------------------------------------

    audience_cache = {}

    existing_audience_result = await db.execute(
        select(TargetAudience)
    )

    existing_audiences = (
        existing_audience_result.scalars().all()
    )

    for audience in existing_audiences:
        audience_cache[
            audience.name.strip().lower()
        ] = audience

    # -----------------------------------------------------
    # PROCESS ROWS
    # -----------------------------------------------------

    for row_number, row in enumerate(
        rows[1:],
        start=2,
    ):

        try:
            recipient_type = normalize_recipient_type(
                get_value(
                    row,
                    "Type",
                )
            )

            name = get_value(
                row,
                "Name",
            )

            email = get_value(
                row,
                "Email",
            )

            whatsapp_number = get_value(
                row,
                "WhatsApp Number",
            )

            target_audience_name = get_value(
                row,
                "Target Audience",
            )

            # -------------------------------------------------
            # IGNORE EMPTY ROW
            # -------------------------------------------------

            if not any(
                [
                    recipient_type,
                    name,
                    email,
                    whatsapp_number,
                    target_audience_name,
                ]
            ):
                continue

            # -------------------------------------------------
            # VALIDATE TYPE
            # -------------------------------------------------

            if recipient_type not in {
                "HOD",
                "LINE_MANAGER",
            }:
                raise ValueError(
                    "Type must be HOD or LINE_MANAGER"
                )

            # -------------------------------------------------
            # VALIDATE NAME
            # -------------------------------------------------

            if not name:
                raise ValueError(
                    "Name is required"
                )

            # -------------------------------------------------
            # VALIDATE EMAIL
            # -------------------------------------------------

            if not email:
                raise ValueError(
                    "Email is required"
                )

            # -------------------------------------------------
            # VALIDATE TARGET AUDIENCE
            # -------------------------------------------------

            if not target_audience_name:
                raise ValueError(
                    "Target Audience is required"
                )

            audience_key = (
                target_audience_name
                .strip()
                .lower()
            )

            audience = audience_cache.get(
                audience_key
            )

            if audience is None:
                audience = TargetAudience(
                    name=target_audience_name.strip()
                )

                db.add(audience)
                await db.flush()

                audience_cache[
                    audience_key
                ] = audience

            # -------------------------------------------------
            # HOD
            # -------------------------------------------------

            if recipient_type == "HOD":

                result = await db.execute(
                    select(HOD).where(
                        func.lower(HOD.email)
                        == email.strip().lower(),
                        HOD.audience_id
                        == audience.id,
                    )
                )

                existing = (
                    result.scalar_one_or_none()
                )

                if existing:

                    existing.name = name.strip()
                    existing.email = email.strip()
                    existing.whatsapp_number = (
                        whatsapp_number
                    )
                    existing.audience_id = (
                        audience.id
                    )

                    updated.append(
                        {
                            "row": row_number,
                            "type": "HOD",
                            "id": existing.id,
                            "name": existing.name,
                            "email": existing.email,
                            "target_audience": (
                                audience.name
                            ),
                            "audience_id": (
                                audience.id
                            ),
                        }
                    )

                else:

                    new_hod = HOD(
                        name=name.strip(),
                        email=email.strip(),
                        whatsapp_number=(
                            whatsapp_number
                        ),
                        audience_id=audience.id,
                        # Legacy department relationship.
                        # Keep it NULL for the new audience-based
                        # workflow.
                        department_id=None,
                    )

                    db.add(new_hod)
                    await db.flush()

                    created.append(
                        {
                            "row": row_number,
                            "type": "HOD",
                            "id": new_hod.id,
                            "name": new_hod.name,
                            "email": new_hod.email,
                            "target_audience": (
                                audience.name
                            ),
                            "audience_id": (
                                audience.id
                            ),
                        }
                    )

            # -------------------------------------------------
            # LINE MANAGER
            # -------------------------------------------------

            else:

                result = await db.execute(
                    select(LineManager).where(
                        func.lower(LineManager.email)
                        == email.strip().lower(),
                        LineManager.audience_id
                        == audience.id,
                    )
                )

                existing = (
                    result.scalar_one_or_none()
                )

                if existing:

                    existing.name = name.strip()
                    existing.email = email.strip()
                    existing.whatsapp_number = (
                        whatsapp_number
                    )
                    existing.audience_id = (
                        audience.id
                    )

                    updated.append(
                        {
                            "row": row_number,
                            "type": "LINE_MANAGER",
                            "id": existing.id,
                            "name": existing.name,
                            "email": existing.email,
                            "target_audience": (
                                audience.name
                            ),
                            "audience_id": (
                                audience.id
                            ),
                        }
                    )

                else:

                    new_manager = LineManager(
                        name=name.strip(),
                        email=email.strip(),
                        whatsapp_number=(
                            whatsapp_number
                        ),
                        audience_id=audience.id,
                        # Legacy department relationship.
                        # Keep it NULL for the new audience-based
                        # workflow.
                        department_id=None,
                    )

                    db.add(new_manager)
                    await db.flush()

                    created.append(
                        {
                            "row": row_number,
                            "type": "LINE_MANAGER",
                            "id": new_manager.id,
                            "name": new_manager.name,
                            "email": new_manager.email,
                            "target_audience": (
                                audience.name
                            ),
                            "audience_id": (
                                audience.id
                            ),
                        }
                    )

        except Exception as exc:

            errors.append(
                {
                    "row": row_number,
                    "error": str(exc),
                }
            )

    # -----------------------------------------------------
    # ROLLBACK IF ANY ERROR EXISTS
    # -----------------------------------------------------

    if errors:

        await db.rollback()

        return {
            "success": False,
            "message": (
                "Import failed because one or more "
                "rows contain errors. No changes "
                "were committed."
            ),
            "summary": {
                "created": 0,
                "updated": 0,
                "errors": len(errors),
                "total_processed": 0,
            },
            "created": [],
            "updated": [],
            "errors": errors,
        }

    # -----------------------------------------------------
    # COMMIT
    # -----------------------------------------------------

    try:
        await db.commit()

    except Exception as exc:
        await db.rollback()

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to commit master recipient "
                f"data: {str(exc)}"
            ),
        )

    # -----------------------------------------------------
    # RESPONSE
    # -----------------------------------------------------

    return {
        "success": True,
        "message": (
            "Master recipient data imported successfully."
        ),
        "summary": {
            "created": len(created),
            "updated": len(updated),
            "errors": 0,
            "total_processed": (
                len(created)
                + len(updated)
            ),
        },
        "created": created,
        "updated": updated,
        "errors": [],
    }
